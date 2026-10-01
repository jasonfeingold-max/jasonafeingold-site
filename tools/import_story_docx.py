#!/usr/bin/env python3
"""Convert a final-draft DOCX manuscript into a verified story page."""

from __future__ import annotations

import argparse
import html
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document


SCENE_BREAK = "***"


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def is_scene_break(text: str) -> bool:
    return re.sub(r"\s+", "", text) == SCENE_BREAK


def render_runs(paragraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        raw_text = run.text.replace("\t", " ")
        text = html.escape(raw_text)
        if not raw_text:
            continue
        if not raw_text.strip():
            parts.append(text)
            continue
        if run.underline:
            text = f"<cite>{text}</cite>"
        elif run.italic:
            text = f"<em>{text}</em>"
        if run.bold:
            text = f"<strong>{text}</strong>"
        parts.append(text)
    return "".join(parts).strip()


def manuscript_segments(docx_path: Path, title: str):
    document = Document(docx_path)
    title_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if normalize(paragraph.text) == normalize(title)
    )

    segments: list[tuple[str, str]] = []
    italics: list[str] = []
    for paragraph in document.paragraphs[title_index + 1 :]:
        plain = normalize(paragraph.text)
        if not plain:
            continue
        if is_scene_break(plain):
            segments.append(("break", SCENE_BREAK))
            continue
        segments.append(("paragraph", plain))
        italics.extend(normalize(run.text) for run in paragraph.runs if run.italic and normalize(run.text))
    return segments, italics


def manuscript_body(docx_path: Path, title: str) -> str:
    document = Document(docx_path)
    title_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if normalize(paragraph.text) == normalize(title)
    )

    body: list[str] = []
    for paragraph in document.paragraphs[title_index + 1 :]:
        plain = normalize(paragraph.text)
        if not plain:
            continue
        if is_scene_break(plain):
            body.append('        <div class="scene-break" aria-label="Scene break">* &nbsp; * &nbsp; *</div>')
            continue
        body.append(f"        <p>{render_runs(paragraph)}</p>")
    return "\n".join(body)


class StoryBodyParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_story_body = False
        self.body_div_depth = 0
        self.in_paragraph = False
        self.in_italic = False
        self.paragraph_parts: list[str] = []
        self.italic_parts: list[str] = []
        self.segments: list[tuple[str, str]] = []
        self.italics: list[str] = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "div" and "story-body" in attrs.get("class", "").split():
            self.in_story_body = True
            self.body_div_depth = 1
            return
        if not self.in_story_body:
            return
        if tag == "div":
            self.body_div_depth += 1
            if "scene-break" in attrs.get("class", "").split():
                self.segments.append(("break", SCENE_BREAK))
        elif tag == "p":
            self.in_paragraph = True
            self.paragraph_parts = []
        elif tag == "em":
            self.in_italic = True
            self.italic_parts = []

    def handle_endtag(self, tag):
        if not self.in_story_body:
            return
        if tag == "p" and self.in_paragraph:
            self.segments.append(("paragraph", normalize("".join(self.paragraph_parts))))
            self.in_paragraph = False
        elif tag == "em" and self.in_italic:
            self.italics.append(normalize("".join(self.italic_parts)))
            self.in_italic = False
        elif tag == "div":
            self.body_div_depth -= 1
            if self.body_div_depth == 0:
                self.in_story_body = False

    def handle_data(self, data):
        if self.in_paragraph:
            self.paragraph_parts.append(data)
        if self.in_italic:
            self.italic_parts.append(data)


def verify_import(docx_path: Path, title: str, page: str) -> dict[str, object]:
    source_segments, source_italics = manuscript_segments(docx_path, title)
    parser = StoryBodyParser()
    parser.feed(page)
    source_breaks = sum(kind == "break" for kind, _ in source_segments)
    imported_breaks = sum(kind == "break" for kind, _ in parser.segments)
    checks = {
        "text_and_paragraphs_match": source_segments == parser.segments,
        "italics_match": source_italics == parser.italics,
        "scene_breaks_match": source_breaks == imported_breaks,
        "source_paragraphs": sum(kind == "paragraph" for kind, _ in source_segments),
        "imported_paragraphs": sum(kind == "paragraph" for kind, _ in parser.segments),
        "source_italic_runs": len(source_italics),
        "imported_italic_runs": len(parser.italics),
        "source_scene_breaks": source_breaks,
        "imported_scene_breaks": imported_breaks,
    }
    checks["passed"] = all(
        checks[key]
        for key in ("text_and_paragraphs_match", "italics_match", "scene_breaks_match")
    )
    return checks


def build_page(args) -> str:
    source_title = args.source_title or args.title
    body = manuscript_body(Path(args.input), source_title)
    title = html.escape(args.title)
    site_author = html.escape(args.site_author)
    byline = html.escape(args.byline or args.site_author)
    publication = html.escape(args.publication)
    year = html.escape(args.year)
    original_link = ""
    if args.archive:
        archive = html.escape(args.archive, quote=True)
        original_link = (
            f' <a href="{archive}" target="_blank" rel="noopener">'
            f'{html.escape(args.archive_label)} <span aria-hidden="true">↗</span></a>'
        )

    return f'''<!doctype html>
<html lang="en">
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <meta name="description" content="{title}, a short story by {byline}." />
    <title>{title} | {site_author}</title>
    <link rel="stylesheet" href="../styles.css" />
  </head>
  <body>
    <header class="site-header story-site-header">
      <a class="wordmark" href="../index.html#top" aria-label="{site_author}, home">
        <span class="key-mark" aria-hidden="true">JAF</span>
        <span class="wordmark-text">{site_author}</span>
      </a>

      <nav aria-label="Primary navigation">
        <a href="../index.html#work">Work</a>
        <a href="../index.html#bibliography">Bibliography</a>
        <a href="../index.html#about">About</a>
      </nav>
    </header>

    <main class="story-main" id="top">
      <article class="story">
        <header class="story-masthead">
          <p class="eyebrow">Short fiction</p>
          <h1>{title}</h1>
          <p class="story-byline">By {byline}</p>
          <p class="story-publication">
            Originally published by <cite>{publication}</cite> in {year}.{original_link}
          </p>
        </header>

        <div class="story-body">
{body}
        </div>

        <div class="story-end">
          <span aria-hidden="true">№ {html.escape(args.story_number)}</span>
          <a class="text-link" href="../index.html#bibliography">Return to bibliography <span aria-hidden="true">←</span></a>
        </div>
      </article>
    </main>

    <footer>
      <p>© <span id="year">2026</span> {site_author}</p>
      <a href="#top">Back to top <span aria-hidden="true">↑</span></a>
    </footer>

    <script>
      document.getElementById("year").textContent = new Date().getFullYear();
    </script>
  </body>
</html>
'''


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--title", required=True)
    parser.add_argument("--source-title")
    parser.add_argument("--site-author", default="Jason A. Feingold")
    parser.add_argument("--byline")
    parser.add_argument("--publication", required=True)
    parser.add_argument("--year", required=True)
    parser.add_argument("--archive")
    parser.add_argument("--archive-label", default="View the original archive")
    parser.add_argument("--story-number", default="17")
    parser.add_argument("--report")
    args = parser.parse_args()

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    page = build_page(args)
    output.write_text(page, encoding="utf-8")
    report = verify_import(Path(args.input), args.source_title or args.title, page)
    report.update({"title": args.title, "source": str(Path(args.input)), "output": str(output)})
    if args.report:
        Path(args.report).write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
